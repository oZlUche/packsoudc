# -*- coding: utf-8 -*-
"""把作者在电子表格里画的地图转成仓库里的纯文本格式（架构文档 5.4.1）。

    python tools/map_xlsx_to_text.py <地图.xlsx> <输出.txt>

电子表格只是顺手的编辑器，仓库里存的是文本。文本分两段，行列一一对齐：

    [map]      符号层：节点、同环通路、异环通路、行末占位
    [danger]   危险层：'.' 常规 / '1' 中危 / '2' 高危 —— 来自单元格底色

底色是作者标区域的方式（表格里刷个底色比记数字容易），文本没有颜色，
所以拆成第二层。两层行数相同、每行等长，逐格对应。

依赖 openpyxl（pip install openpyxl）。
"""
import io
import sys

import openpyxl

# 底色 → 危险等级。表格里红 = 高危、黄 = 中危、没底色 = 常规
FILL_TO_DANGER = {
    "FFFF0000": "2",
    "FFFFFF00": "1",
}

# 图例所在的列（AL/AM），不属于地图
LEGEND_FIRST_COL = 38


def cell_text(value):
    if value is None:
        return "."
    s = str(value).strip()
    return s if s else "."


def convert(xlsx_path, out_path):
    ws = openpyxl.load_workbook(xlsx_path).worksheets[0]
    last_col = min(ws.max_column, LEGEND_FIRST_COL - 1)

    map_rows, danger_rows = [], []
    for r in range(1, ws.max_row + 1):
        sym, dng = [], []
        for c in range(1, last_col + 1):
            cell = ws.cell(r, c)
            sym.append(cell_text(cell.value))
            fill = cell.fill.fgColor.rgb if (cell.fill is not None and cell.fill.fill_type) else None
            dng.append(FILL_TO_DANGER.get(fill, ".") if cell.value is not None else ".")
        map_rows.append("".join(sym))
        danger_rows.append("".join(dng))

    # 去掉表头行（第 1 行只有图例）与全空的尾行；两层同步裁剪
    while map_rows and set(map_rows[0]) <= {"."}:
        map_rows.pop(0)
        danger_rows.pop(0)
    while map_rows and set(map_rows[-1]) <= {"."}:
        map_rows.pop()
        danger_rows.pop()
    width = max(len(row.rstrip(".")) for row in map_rows)
    map_rows = [row[:width].ljust(width, ".") for row in map_rows]
    danger_rows = [row[:width].ljust(width, ".") for row in danger_rows]

    text = "[map]\n" + "\n".join(map_rows) + "\n\n[danger]\n" + "\n".join(danger_rows) + "\n"
    io.open(out_path, "w", encoding="utf-8", newline="\n").write(text)
    return text


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    sys.stdout.reconfigure(encoding="utf-8")
    print(convert(sys.argv[1], sys.argv[2]))
